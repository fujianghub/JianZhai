"""Binary-format documents (PDF / PPT / EPUB / image) in every export format.

Before 2026-09-08 these exported as a bare title with an empty body — the
original file never left the site. Now ``doc_export_body`` synthesizes a
Markdown stand-in (original-file link + rendered slides + notes) and the
existing media pipeline bundles / embeds the files.
"""
from __future__ import annotations

import io
import zipfile

import pytest
from django.core.files.base import ContentFile
from docx import Document as DocxDocument

from apps.editor.models import Attachment, SlideImage
from apps.exporter.scope import collect_for_scope
from apps.exporter.services import common, docx_export, html_export, markdown_export, static_site
from apps.exporter.tests.conftest import make_doc

PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _jpeg(color=(200, 30, 30), size=(64, 40)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="JPEG")
    return buf.getvalue()


def _attach(doc, owner, name: str, data: bytes, mime: str, kind=Attachment.KIND_DOCUMENT):
    return Attachment.objects.create(
        document=doc, uploaded_by=owner, file=ContentFile(data, name=name),
        original_filename=name, kind=kind, mime_type=mime, size=len(data),
    )


@pytest.fixture
def pdf_doc(owner, kb):
    doc = make_doc(kb, "manual", status="published")
    doc.visibility = "public"
    doc.save(update_fields=["visibility"])
    _attach(doc, owner, "手册.pdf", b"%PDF-1.4 fake pdf bytes", "application/pdf")
    return doc


@pytest.fixture
def pptx_doc(owner, kb):
    doc = make_doc(kb, "deck", status="published")
    doc.visibility = "public"
    doc.save(update_fields=["visibility"])
    att = _attach(doc, owner, "deck.pptx", b"PK fake", PPTX_CT)
    for i, notes in enumerate(["第一页备注", "", "第三页备注\n第二行"]):
        SlideImage.objects.create(
            document=doc, source=att, index=i, width=64, height=40, notes=notes,
            image=ContentFile(_jpeg(), name=f"slide-{i}.jpg"),
        )
    return doc


@pytest.fixture
def image_doc(owner, kb):
    doc = make_doc(kb, "photo", status="published")
    doc.visibility = "public"
    doc.save(update_fields=["visibility"])
    _attach(doc, owner, "photo.jpg", _jpeg((10, 200, 10)), "image/jpeg", kind=Attachment.KIND_IMAGE)
    return doc


@pytest.mark.django_db
def test_doc_export_body_synthesizes_binary_markdown(pdf_doc, pptx_doc, image_doc, settings):
    settings.SITE_PUBLIC_URL = "https://jz.example"
    body = common.doc_export_body(pdf_doc)
    assert "PDF 原件" in body and "手册.pdf" in body
    assert '<a href="/media/uploads/' in body
    assert "https://jz.example/media/uploads/" in body  # plain-text fallback
    deck = common.doc_export_body(pptx_doc)
    assert deck.count("![第 ") == 3
    assert "> 备注：第一页备注" in deck and "> 第二行" in deck
    assert "![photo.jpg](/media/uploads/" in common.doc_export_body(image_doc)
    # Text documents are untouched.
    text = make_doc(pdf_doc.knowledge_base, "t", published="# hi", status="published")
    assert common.doc_export_body(text) == "# hi"


@pytest.mark.django_db
def test_markdown_single_pdf_bundles_original(owner, kb, pdf_doc):
    scope = collect_for_scope(owner=owner, scope="doc", target_id=pdf_doc.id)
    path, filename, mime = markdown_export.export(scope)
    assert filename.endswith(".zip") and mime == "application/zip"
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        assert "content.md" in names
        originals = [n for n in names if n.startswith("assets/uploads/") and n.endswith(".pdf")]
        assert len(originals) == 1
        md = zf.read("content.md").decode()
        assert f'href="{originals[0]}"' in md  # rewritten to the zip-relative path
        assert zf.read(originals[0]).startswith(b"%PDF")


@pytest.mark.django_db
def test_markdown_kb_export_ships_slides_and_original(owner, kb, pptx_doc):
    make_doc(kb, "text", published="hello", status="published")
    scope = collect_for_scope(owner=owner, scope="kb", target_id=kb.id)
    path, *_ = markdown_export.export(scope)
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        assert sum(1 for n in names if n.startswith("assets/slides/") and n.endswith(".jpg")) == 3
        assert any(n.startswith("assets/uploads/") and n.endswith(".pptx") for n in names)
        deck_md = next(zf.read(n).decode() for n in names if n.endswith("deck.md"))
        assert "![第 1 页](assets/slides/" in deck_md
        assert "备注：第一页备注" in deck_md


@pytest.mark.django_db
def test_markdown_oversized_original_stays_as_site_link(owner, kb, pdf_doc, settings):
    settings.EXPORT_MAX_ASSET_BYTES = 4
    scope = collect_for_scope(owner=owner, scope="doc", target_id=pdf_doc.id)
    path, filename, _ = markdown_export.export(scope)
    assert filename.endswith(".md")  # nothing bundleable → plain .md
    text = path.read_text()
    assert "文件超过离线包上限" in text and 'href="/media/uploads/' in text


@pytest.mark.django_db
def test_html_anthology_embeds_slides_as_data_uris(owner, kb, pptx_doc, pdf_doc):
    scope = collect_for_scope(owner=owner, scope="kb", target_id=kb.id)
    html = html_export.render_html(scope, mode="print")
    assert html.count('src="data:image/jpeg;base64,') == 3
    assert "PDF 原件" in html and "手册.pdf" in html
    # Small originals ride along as data URIs too (≤ MAX_EMBED_BYTES).
    assert 'href="data:application/pdf;base64,' in html


@pytest.mark.django_db
def test_docx_export_inlines_slides_and_placeholder(owner, kb, pptx_doc, pdf_doc):
    scope = collect_for_scope(owner=owner, scope="kb", target_id=kb.id)
    path, *_ = docx_export.export(scope)
    document = DocxDocument(str(path))
    assert len(document.inline_shapes) == 3
    text = "\n".join(p.text for p in document.paragraphs)
    assert "备注：第一页备注" in text
    assert "PDF 原件" in text and "手册.pdf" in text


@pytest.mark.django_db
def test_docx_follows_kb_tree_order(owner, kb, folder):
    # Folder doc sorts after root docs in the tree (mirrors HTML anthology).
    make_doc(kb, "in-folder", published="x", status="published", folder=folder, order=0)
    make_doc(kb, "root-doc", published="y", status="published", order=5)
    scope = collect_for_scope(owner=owner, scope="kb", target_id=kb.id)
    path, *_ = docx_export.export(scope)
    titles = [p.text for p in DocxDocument(str(path)).paragraphs if p.style.name == "Heading 1"]
    assert titles.index("root-doc") < titles.index("in-folder")


@pytest.mark.django_db
def test_static_site_includes_binary_docs_with_assets(owner, kb, pdf_doc, pptx_doc):
    scope = collect_for_scope(owner=owner, scope="kb", target_id=kb.id, only_published=True)
    path, *_ = static_site.export(scope)
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        pages = [n for n in names if n.endswith(".html") and n != "index.html"]
        assert len(pages) == 2
        assert any(n.startswith("assets/uploads/") and n.endswith(".pdf") for n in names)
        assert sum(1 for n in names if n.startswith("assets/slides/")) == 3
        deck_page = zf.read(next(n for n in pages if n.startswith("deck-"))).decode()
        assert 'src="assets/slides/' in deck_page
        index = zf.read("index.json").decode()
        assert "PDF 原件" in index
